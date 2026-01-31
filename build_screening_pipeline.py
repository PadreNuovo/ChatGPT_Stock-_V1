import argparse
import json
import logging
import os
from dataclasses import dataclass, field
from datetime import datetime, date, timedelta
from pathlib import Path
import re
import sys
import time
import unicodedata
from zoneinfo import ZoneInfo
from typing import Dict, List, Optional, Tuple, Set, Any

import pandas as pd
import requests
from bs4 import BeautifulSoup
from dateutil import parser as dateparser
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import difflib


# --------------------------------------------------------------------------------------
# Data structures
# --------------------------------------------------------------------------------------

@dataclass
class ExternalFetchResult:
    status: str
    message: str
    fetched_at: datetime
    source: str
    records: int = 0
    warnings: List[str] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)


# --------------------------------------------------------------------------------------
# Utility: Column resolution with normalization
# --------------------------------------------------------------------------------------

class ColumnResolver:
    def __init__(self):
        self.audit_logs: List[Dict[str, str]] = []

    @staticmethod
    def normalize_text(text: str) -> str:
        if text is None:
            return ""
        # Normalize width, remove spaces and punctuation like dots, underscores, hyphens
        normalized = unicodedata.normalize("NFKC", str(text))
        normalized = re.sub(r"[\s\-_\.]+", "", normalized)
        normalized = normalized.strip().lower()
        return normalized

    def resolve(self, df: pd.DataFrame, field: str, candidates: List[str], sheet: str) -> str:
        available = list(df.columns)
        cand_norm = [self.normalize_text(c) for c in candidates]
        col_map = {col: self.normalize_text(col) for col in available}

        # 1) exact match
        for col in available:
            if col in candidates:
                self.audit_logs.append(
                    {
                        "field_name": field,
                        "sheet": sheet,
                        "resolved_column": col,
                        "method": "exact",
                        "confidence": "1.0",
                    }
                )
                return col

        # 2) case-insensitive
        for col in available:
            if col.lower() in [c.lower() for c in candidates]:
                self.audit_logs.append(
                    {
                        "field_name": field,
                        "sheet": sheet,
                        "resolved_column": col,
                        "method": "casefold",
                        "confidence": "0.9",
                    }
                )
                return col

        # 3) normalized
        for col in available:
            if col_map[col] in cand_norm:
                self.audit_logs.append(
                    {
                        "field_name": field,
                        "sheet": sheet,
                        "resolved_column": col,
                        "method": "normalized",
                        "confidence": "0.8",
                    }
                )
                return col

        # 4) close matches using normalized tokens
        normalized_available = list(col_map.values())
        best = None
        for cand in cand_norm:
            matches = difflib.get_close_matches(cand, normalized_available, n=1, cutoff=0.6)
            if matches:
                best = matches[0]
                break
        if best:
            # find original column with matching normalized
            for col, norm in col_map.items():
                if norm == best:
                    self.audit_logs.append(
                        {
                            "field_name": field,
                            "sheet": sheet,
                            "resolved_column": col,
                            "method": "fuzzy",
                            "confidence": "0.6",
                        }
                    )
                    # Although fuzzy found, spec requires stopping with error and suggestion
                    raise ValueError(
                        f"Required field '{field}' not found in sheet '{sheet}'. Closest match: '{col}'."
                    )

        raise ValueError(
            f"Required field '{field}' not found in sheet '{sheet}'. Candidates: {candidates}. Columns: {available}"
        )


# --------------------------------------------------------------------------------------
# Excel IO helper
# --------------------------------------------------------------------------------------

class ExcelIO:
    def __init__(self, path: Path):
        self.path = path
        self.wb = load_workbook(path)

    def read_sheet(self, name: str) -> pd.DataFrame:
        if name not in self.wb.sheetnames:
            raise ValueError(f"Missing required sheet: {name}")
        df = pd.read_excel(self.path, sheet_name=name, engine="openpyxl")
        return df

    def replace_sheet_with_df(self, name: str, df: pd.DataFrame):
        if name in self.wb.sheetnames:
            idx = self.wb.sheetnames.index(name)
            self.wb.remove(self.wb[name])
            ws = self.wb.create_sheet(name, idx)
        else:
            ws = self.wb.create_sheet(name)
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in df.itertuples(index=False):
            ws.append(list(row))
        for col_idx, _ in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 15
        ws.auto_filter.ref = ws.dimensions

    def save(self, out_path: Path):
        self.wb.save(out_path)


# --------------------------------------------------------------------------------------
# TSE PBR list normalizer
# --------------------------------------------------------------------------------------

class TSEListNormalizer:
    REQUIRED_COLS = [
        "Code4",
        "CompanyName",
        "Market",
        "Status",
        "UpdateDate",
        "EnglishFlag",
        "InvestorContactFlag",
    ]

    def __init__(self, session: Optional[requests.Session] = None):
        self.session = session or requests.Session()

    def fetch_latest_xlsx_url(self, page_url: str, override_url: Optional[str]) -> str:
        if override_url:
            return override_url
        resp = self.session.get(page_url, timeout=30)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        links = soup.find_all("a")
        xlsx_links = [link.get("href") for link in links if link.get("href") and "xlsx" in link.get("href")]
        if not xlsx_links:
            raise RuntimeError("No XLSX link found on TSE page")
        href = xlsx_links[0]
        if href.startswith("http"):
            return href
        # relative link
        base = page_url.rsplit("/", 1)[0]
        return base + "/" + href.lstrip("./")

    def download_xlsx(self, url: str, cache_dir: Path, refresh: bool) -> Path:
        cache_dir.mkdir(parents=True, exist_ok=True)
        fname = cache_dir / (re.sub(r"[^A-Za-z0-9]+", "_", url) + ".xlsx")
        if fname.exists() and not refresh:
            return fname
        resp = self.session.get(url, timeout=60)
        resp.raise_for_status()
        with open(fname, "wb") as f:
            f.write(resp.content)
        return fname

    def load_workbook_to_raw_df(self, xlsx_path: Path) -> pd.DataFrame:
        wb = load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        data = []
        max_row = ws.max_row
        max_col = ws.max_column
        # forward fill merged cell values
        grid = [[None for _ in range(max_col)] for _ in range(max_row)]
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                val = ws.cell(r, c).value
                grid[r - 1][c - 1] = val
        # forward fill horizontally and vertically for merged-like structures
        for r in range(max_row):
            for c in range(max_col):
                if grid[r][c] is None:
                    if r > 0 and grid[r - 1][c] is not None:
                        grid[r][c] = grid[r - 1][c]
                    elif c > 0 and grid[r][c - 1] is not None:
                        grid[r][c] = grid[r][c - 1]
        for row in grid:
            data.append(row)
        df = pd.DataFrame(data)
        return df

    def normalize(self, raw_df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict]:
        audit = {}
        # attempt to find header row by locating column containing code pattern
        header_row_idx = None
        for i, row in raw_df.iterrows():
            row_str = row.astype(str).tolist()
            joined = " ".join(row_str)
            if re.search(r"コード|証券コード|Code", joined):
                header_row_idx = i
                break
        if header_row_idx is None:
            header_row_idx = 0
        header = raw_df.iloc[header_row_idx].tolist()
        df = raw_df.iloc[header_row_idx + 1 :].copy()
        df.columns = header
        df = df.dropna(how="all")
        df = df.rename(columns={col: str(col).strip() for col in df.columns})

        resolver = ColumnResolver()
        required = {
            "Code4": ["コード", "証券コード", "Code", "code"],
            "CompanyName": ["会社名", "銘柄名", "Name", "Company"],
            "Market": ["市場", "市場区分", "Market"],
            "Status": ["対応状況", "ステータス", "Status"],
            "UpdateDate": ["更新日", "Update", "更新日時"],
            "EnglishFlag": ["英語情報", "English", "英語"],
            "InvestorContactFlag": ["投資家窓口", "Contact", "窓口"],
        }
        cols = {}
        for field, cand in required.items():
            try:
                cols[field] = resolver.resolve(df, field, cand, "TSE_PBR_List")
            except ValueError:
                cols[field] = None
        norm_df = pd.DataFrame(columns=self.REQUIRED_COLS)
        if any(v is None for v in cols.values()):
            audit["column_resolution_warnings"] = "Some columns missing; returned empty normalized frame"
            return norm_df, audit

        norm_df["Code4"] = df[cols["Code4"]].apply(lambda x: self._normalize_code(str(x)))
        norm_df["CompanyName"] = df[cols["CompanyName"]]
        norm_df["Market"] = df[cols["Market"]]
        norm_df["Status"] = df[cols["Status"]].apply(self._normalize_status)
        norm_df["UpdateDate"] = df[cols["UpdateDate"]].apply(self._to_date_str)
        norm_df["EnglishFlag"] = df[cols["EnglishFlag"]].apply(self._to_flag)
        norm_df["InvestorContactFlag"] = df[cols["InvestorContactFlag"]].apply(self._to_flag)

        audit["column_resolution"] = resolver.audit_logs
        return norm_df, audit

    @staticmethod
    def _normalize_code(val: str) -> Optional[str]:
        if val is None or pd.isna(val):
            return pd.NA
        digits = re.findall(r"\d{4}", str(val))
        if digits:
            return digits[0]
        return pd.NA

    @staticmethod
    def _normalize_status(val: str) -> Optional[str]:
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return pd.NA
        text = str(val)
        if any(k in text for k in ["開示", "公表", "済"]):
            return "開示済"
        if any(k in text for k in ["検討", "対応中", "準備"]):
            return "検討中"
        if any(k in text for k in ["その他", "未定"]):
            return "その他"
        return pd.NA

    @staticmethod
    def _to_flag(val) -> Optional[int]:
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return pd.NA
        text = str(val).strip().lower()
        truthy = {"1", "yes", "y", "true", "〇", "○", "有", "あり", "✔", "✓"}
        falsy = {"0", "no", "false", "無", "なし"}
        if text in truthy:
            return 1
        if text in falsy:
            return 0
        return pd.NA

    @staticmethod
    def _to_date_str(val) -> Optional[str]:
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return pd.NA
        try:
            dt = dateparser.parse(str(val))
            return dt.date().isoformat()
        except Exception:
            return pd.NA

    def run(self, page_url: str, override_url: Optional[str], cache_dir: Path, refresh: bool) -> Tuple[pd.DataFrame, ExternalFetchResult, Dict]:
        start = datetime.now()
        try:
            xlsx_url = self.fetch_latest_xlsx_url(page_url, override_url)
            path = self.download_xlsx(xlsx_url, cache_dir, refresh)
            raw = self.load_workbook_to_raw_df(path)
            norm, audit = self.normalize(raw)
            result = ExternalFetchResult(
                status="OK",
                message="Fetched and normalized",
                fetched_at=datetime.now(),
                source="TSE",
                records=len(norm),
            )
            audit.update({"url": xlsx_url, "path": str(path), "code_unique": norm["Code4"].nunique() if not norm.empty else 0})
            if norm.empty:
                result.status = "FAILED"
                result.message = "Normalization failed"
            return norm, result, audit
        except Exception as exc:  # noqa: BLE001
            logging.exception("TSE fetch failed")
            empty_df = pd.DataFrame(columns=self.REQUIRED_COLS)
            result = ExternalFetchResult(
                status="FAILED",
                message=str(exc),
                fetched_at=datetime.now(),
                source="TSE",
                records=0,
                errors=[str(exc)],
            )
            return empty_df, result, {"error": str(exc)}


# --------------------------------------------------------------------------------------
# EDINET large holding extractor
# --------------------------------------------------------------------------------------

class EdinetLargeHoldingExtractor:
    REQUIRED_COLS = [
        "docID",
        "submitDateTime",
        "fileDate",
        "secCode",
        "Code4",
        "docDescription",
        "formCode",
        "docTypeCode",
        "MatchRule",
        "IsWeakMatch",
        "IsActive",
        "EventType",
    ]

    def __init__(self, session: Optional[requests.Session] = None):
        self.session = session or requests.Session()

    def load_form_code_list(self, url: str, cache_dir: Path, refresh: bool) -> Set[Tuple[str, str]]:
        cache_dir.mkdir(parents=True, exist_ok=True)
        fname = cache_dir / "edinet_form_codes.xlsx"
        if fname.exists() and not refresh:
            df = pd.read_excel(fname, engine="openpyxl")
        else:
            resp = self.session.get(url, timeout=60)
            resp.raise_for_status()
            with open(fname, "wb") as f:
                f.write(resp.content)
            df = pd.read_excel(fname, engine="openpyxl")
        df = df.rename(columns={col: str(col) for col in df.columns})
        df = df.fillna("")
        allowed = set()
        for _, row in df.iterrows():
            name = str(row.iloc[2]) if len(row) > 2 else ""
            normalized_name = ColumnResolver.normalize_text(name)
            if any(k in normalized_name for k in ["大量保有報告書", "変更報告書"]):
                form = str(row.iloc[0]) if len(row) > 0 else ""
                dtype = str(row.iloc[1]) if len(row) > 1 else ""
                allowed.add((form.strip(), dtype.strip()))
        return allowed

    def fetch_doclist(self, d: date, api_key: str, cache_dir: Path, refresh: bool) -> dict:
        cache_dir.mkdir(parents=True, exist_ok=True)
        fname = cache_dir / f"edinet_doclist_{d.isoformat()}.json"
        if fname.exists() and not refresh:
            with open(fname, "r", encoding="utf-8") as f:
                return json.load(f)
        params = {"date": d.isoformat(), "type": 2}
        headers = {"X-API-KEY": api_key} if api_key else {}
        delay = 2
        for attempt in range(5):
            resp = self.session.get(
                "https://disclosure.edinet-fsa.go.jp/api/v2/documents.json",
                params=params,
                headers=headers,
                timeout=60,
            )
            if resp.status_code == 200:
                data = resp.json()
                with open(fname, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False)
                return data
            if resp.status_code in {429, 500, 503}:
                time.sleep(delay)
                delay *= 2
                continue
            resp.raise_for_status()
        raise RuntimeError(f"EDINET doclist fetch failed for {d}")

    def classify_and_normalize(self, records: List[dict], allowed_pairs: Set[Tuple[str, str]], fallback_keywords: List[str]) -> pd.DataFrame:
        rows = []
        for rec in records:
            form = str(rec.get("formCode", ""))
            dtype = str(rec.get("docTypeCode", ""))
            desc = str(rec.get("docDescription", ""))
            pair = (form, dtype)
            match_rule = None
            weak = False
            if pair in allowed_pairs:
                match_rule = "A"
            else:
                normalized_desc = ColumnResolver.normalize_text(desc)
                if any(k in normalized_desc for k in fallback_keywords):
                    match_rule = "B"
                    weak = True
            if match_rule is None:
                continue
            # Determine status fields
            withdrawal = str(rec.get("withdrawalStatus", "")).lower()
            disclose = str(rec.get("disclosureStatus", "")).lower()
            edit_status = str(rec.get("docInfoEditStatus", "")).lower()
            is_active = not ("withdraw" in withdrawal or "取消" in withdrawal)
            event_type = "SUBMIT" if "submit" in disclose or disclose == "" else disclose.upper()

            sec_code = rec.get("secCode")
            code4 = pd.NA
            if sec_code:
                digits = re.findall(r"\d{4}", str(sec_code))
                if digits:
                    code4 = digits[0]
            row = {
                "docID": rec.get("docID"),
                "submitDateTime": rec.get("submitDateTime"),
                "fileDate": rec.get("docInfoEditStatusDate"),
                "secCode": sec_code,
                "Code4": code4,
                "docDescription": desc,
                "formCode": form,
                "docTypeCode": dtype,
                "MatchRule": match_rule,
                "IsWeakMatch": weak,
                "IsActive": is_active,
                "EventType": event_type,
            }
            rows.append(row)
        df = pd.DataFrame(rows, columns=self.REQUIRED_COLS)
        return df

    def build_snapshot(self, df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df.copy()
        df = df.copy()
        df["submitDateTime_parsed"] = pd.to_datetime(df["submitDateTime"], errors="coerce")
        df["fileDate_parsed"] = pd.to_datetime(df["fileDate"], errors="coerce")
        df = df.sort_values(["docID", "submitDateTime_parsed", "fileDate_parsed"], ascending=[True, False, False])
        latest = df.groupby("docID").head(1)
        latest = latest.drop(columns=["submitDateTime_parsed", "fileDate_parsed"])
        return latest

    def run(
        self,
        lookback_days: int,
        form_code_url: str,
        api_key: str,
        cache_dir: Path,
        refresh: bool,
    ) -> Tuple[pd.DataFrame, ExternalFetchResult, Dict]:
        start = datetime.now()
        audit = {"lookback_days": lookback_days}
        try:
            allowed = self.load_form_code_list(form_code_url, cache_dir, refresh)
            audit["allowed_pairs"] = len(allowed)
            all_records = []
            for i in range(lookback_days):
                d = date.today() - timedelta(days=i)
                doclist = self.fetch_doclist(d, api_key, cache_dir, refresh)
                docs = doclist.get("results", []) if isinstance(doclist, dict) else []
                all_records.extend(docs)
            audit["total_records"] = len(all_records)
            df = self.classify_and_normalize(all_records, allowed, ["大量保有", "変更報告"])
            snapshot = self.build_snapshot(df)
            audit["match_A"] = (df["MatchRule"] == "A").sum()
            audit["match_B"] = (df["MatchRule"] == "B").sum()
            audit["active_count"] = (df["IsActive"] == True).sum()  # noqa: E712
            audit["secCode_missing"] = df["Code4"].isna().sum()
            result = ExternalFetchResult(
                status="OK",
                message="Fetched",
                fetched_at=datetime.now(),
                source="EDINET",
                records=len(snapshot),
            )
            if snapshot.empty:
                result.status = "FAILED"
                result.message = "No records matched"
            return snapshot, result, audit
        except Exception as exc:  # noqa: BLE001
            logging.exception("EDINET fetch failed")
            empty_df = pd.DataFrame(columns=self.REQUIRED_COLS)
            result = ExternalFetchResult(
                status="FAILED",
                message=str(exc),
                fetched_at=datetime.now(),
                source="EDINET",
                records=0,
                errors=[str(exc)],
            )
            return empty_df, result, {"error": str(exc)}


# --------------------------------------------------------------------------------------
# Screening pipeline
# --------------------------------------------------------------------------------------

class ScreeningPipeline:
    REQUIRED_SHEETS = ["Universe", "Market", "Fundamentals"]

    def __init__(
        self,
        workbook: Path,
        outdir: Path,
        no_external: bool,
        refresh_cache: bool,
        log_level: str,
        watchlist_size: int,
        max_history_runs: int,
        report_only: bool,
        snapshot_bucket_mode: str,
    ):
        self.workbook = workbook
        self.outdir = outdir
        self.no_external = no_external
        self.refresh_cache = refresh_cache
        self.log_level = log_level
        self.watchlist_size = watchlist_size
        self.max_history_runs = max_history_runs
        self.report_only = report_only
        self.snapshot_bucket_mode = snapshot_bucket_mode
        self.io = ExcelIO(workbook)
        self.resolver = ColumnResolver()
        self.audit: Dict[str, Any] = {"run_at": datetime.now().isoformat()}
        self.params: Dict[str, Any] = {}
        self.cache_dir: Path = Path("./data_cache")
        self.external_results: Dict[str, ExternalFetchResult] = {}
        self.external_data: Dict[str, pd.DataFrame] = {}
        self.code_warnings: List[str] = []

    # ------------------------ input loading ------------------------
    def load_inputs(self) -> Dict[str, pd.DataFrame]:
        dfs = {}
        for sheet in self.REQUIRED_SHEETS:
            dfs[sheet] = self.io.read_sheet(sheet)
        # Validate required columns
        universe_fields = {
            "Code": ["Code", "コード", "証券コード"],
            "Name": ["Name", "会社名", "銘柄名"],
            "MarketSegment": ["MarketSegment", "市場", "市場区分"],
            "Sector33": ["Sector33", "33業種", "業種"],
        }
        market_fields = {
            "Code": ["Code", "コード", "証券コード"],
            "MktCap": ["MktCap", "時価総額"],
            "PBR": ["PBR", "株価純資産倍率"],
            "ADV20Value": ["ADV20Value", "20日平均売買代金", "売買代金20日"],
        }
        fundamentals_fields = {
            "Code": ["Code", "コード", "証券コード"],
            "Equity": ["Equity", "純資産", "自己資本"],
            "TotalAssets": ["TotalAssets", "総資産"],
            "Cash": ["Cash", "現金預金"],
            "STInvest": ["STInvest", "短期有価証券", "短期投資"],
            "InterestDebt": ["InterestDebt", "有利子負債"],
            "RetainedEarnings": ["RetainedEarnings", "利益剰余金"],
        }
        # Resolve columns
        resolved = {}
        resolved["Universe"] = {f: self.resolver.resolve(dfs["Universe"], f, cand, "Universe") for f, cand in universe_fields.items()}
        resolved["Market"] = {f: self.resolver.resolve(dfs["Market"], f, cand, "Market") for f, cand in market_fields.items()}
        resolved["Fundamentals"] = {f: self.resolver.resolve(dfs["Fundamentals"], f, cand, "Fundamentals") for f, cand in fundamentals_fields.items()}
        self.audit["column_resolution"] = self.resolver.audit_logs

        # Normalize Code
        for key, mapping in resolved.items():
            df = dfs[key].copy()
            code_col = mapping["Code"]
            df["Code"] = df[code_col].apply(lambda v: self.normalize_code(v, origin_sheet=key))
            for fld, col in mapping.items():
                if fld != "Code":
                    df[fld] = df[col]
            dfs[key] = df[list(mapping.keys())]
        if self.code_warnings:
            self.audit["code_warnings"] = self.code_warnings
        return dfs

    def normalize_code(self, val, origin_sheet: str = ""):
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return pd.NA
        text = str(val)
        if text.isdigit():
            return text.zfill(4)
        digits = re.findall(r"\d{4}", text)
        if digits:
            return digits[0]
        warning = f"Unnormalized code '{text}' in {origin_sheet}"
        logging.warning(warning)
        self.code_warnings.append(warning)
        return text

    # ------------------------ params ------------------------
    def load_params(self):
        default_params = {
            "ADV20Value_min": 50_000_000,
            "PBR_threshold": 1.0,
            "LOOKBACK_DAYS_EDINET": 14,
            "WEIGHT_VALUE": 30,
            "WEIGHT_BALANCE": 35,
            "WEIGHT_LIQUIDITY": 20,
            "WEIGHT_CATALYST": 15,
            "PBR_good_floor": 0.50,
            "NETCASHRATIO_CAP": 0.30,
            "RETEARNRATIO_CAP": 0.60,
            "CATALYST_TSE_DISCLOSED_POINTS": 8,
            "CATALYST_TSE_ENG_POINTS": 2,
            "CATALYST_TSE_CONTACT_POINTS": 1,
            "CATALYST_EDINET_NEW_POINTS": 7,
            "CATALYST_EDINET_MULTIPLIER": 1.0,
            "CATALYST_MAX": 15,
            "SCORE_MODE": "STRICT",
            "SCORE_MOVE_THRESHOLD": 5,
            "RANK_MOVE_THRESHOLD": 10,
            "TSE_PBR_PAGE_URL": "https://www.jpx.co.jp/equities/follow-up/02.html",
            "TSE_PBR_LIST_XLSX_URL": "",
            "EDINET_FORM_CODE_LIST_URL": "https://disclosure2dl.edinet-fsa.go.jp/guide/static/disclosure/download/ESE140327.xlsx",
            "CACHE_DIR": "./data_cache",
        }
        if "Params" in self.io.wb.sheetnames:
            df = pd.read_excel(self.workbook, sheet_name="Params", engine="openpyxl")
            for _, row in df.iterrows():
                if len(row) >= 2:
                    key = row.iloc[0]
                    val = row.iloc[1]
                    default_params[key] = val
        env_tse_url = os.environ.get("TSE_PBR_LIST_XLSX_URL")
        if env_tse_url:
            default_params["TSE_PBR_LIST_XLSX_URL"] = env_tse_url
        self.params = default_params
        self.cache_dir = Path(str(self.params.get("CACHE_DIR", "./data_cache")))
        self.audit["params"] = self.params

    @staticmethod
    def _run_id_now() -> str:
        return datetime.now(tz=ZoneInfo("Asia/Tokyo")).isoformat()

    @staticmethod
    def _score_column(params: Dict[str, Any], df: pd.DataFrame) -> str:
        if params.get("SCORE_MODE", "STRICT") != "STRICT" and "TotalScoreAdj" in df.columns:
            return "TotalScoreAdj"
        return "TotalScore"

    @staticmethod
    def _add_missing_fields(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        core_fields = ["PBR_used", "NetCashRatio", "RetainedEarningsRatio", "ADV20Value"]
        available_fields = [col for col in core_fields if col in df.columns]
        if not available_fields:
            df["MissingCount"] = 0
            df["MissingFields"] = ""
            return df
        missing_mask = df[available_fields].isna()
        df["MissingCount"] = missing_mask.sum(axis=1)
        df["MissingFields"] = missing_mask.apply(
            lambda r: ",".join([col for col, missing in r.items() if missing]),
            axis=1,
        )
        return df

    @staticmethod
    def _normalize_code_series(series: pd.Series) -> pd.Series:
        def normalize(val):
            if val is None or (isinstance(val, float) and pd.isna(val)) or val is pd.NA:
                return pd.NA
            text = str(val)
            return text.zfill(4) if text.isdigit() else text
        return series.apply(normalize)

    def _load_snapshot_history(self, out_path: Path) -> pd.DataFrame:
        if "Snapshot_History" in self.io.wb.sheetnames:
            return pd.read_excel(self.workbook, sheet_name="Snapshot_History", engine="openpyxl")
        if out_path.exists():
            try:
                return pd.read_excel(out_path, sheet_name="Snapshot_History", engine="openpyxl")
            except ValueError:
                return pd.DataFrame()
        return pd.DataFrame()

    def _trim_snapshot_history(self, history: pd.DataFrame) -> pd.DataFrame:
        if history.empty or self.max_history_runs <= 0 or "run_id" not in history.columns:
            return history
        history = history.copy()
        history["run_id_parsed"] = pd.to_datetime(history["run_id"], errors="coerce")
        runs = (
            history[["run_id", "run_id_parsed"]]
            .dropna(subset=["run_id_parsed"])
            .drop_duplicates()
            .sort_values("run_id_parsed")
        )
        if len(runs) <= self.max_history_runs:
            return history.drop(columns=["run_id_parsed"])
        keep_runs = runs["run_id"].iloc[-self.max_history_runs :].tolist()
        trimmed = history[history["run_id"].isin(keep_runs)].drop(columns=["run_id_parsed"])
        return trimmed

    def _select_watchlist(self, scores: pd.DataFrame) -> pd.DataFrame:
        df = self._add_missing_fields(scores)
        df["Code"] = self._normalize_code_series(df["Code"])
        score_col = self._score_column(self.params, df)
        confirmed = df[df["BaseScreenFlag"] == True].sort_values(score_col, ascending=False)  # noqa: E712
        watchlist = confirmed.head(self.watchlist_size).copy()
        watchlist["Bucket"] = "Confirmed"
        remaining_slots = self.watchlist_size - len(watchlist)
        if remaining_slots > 0:
            unconfirmed_pool = df[df["BaseScreenFlag"].isna()].copy()
            unconfirmed_pool = unconfirmed_pool.sort_values(["MissingCount", score_col], ascending=[True, False])
            unconfirmed = unconfirmed_pool.head(remaining_slots).copy()
            unconfirmed["Bucket"] = "Unconfirmed"
            watchlist = pd.concat([watchlist, unconfirmed], ignore_index=True)
        remaining_slots = self.watchlist_size - len(watchlist)
        if remaining_slots > 0:
            supplement_pool = df[df["BaseScreenFlag"] == False].copy()  # noqa: E712
            supplement_pool = supplement_pool.sort_values(score_col, ascending=False)
            supplement = supplement_pool.head(remaining_slots).copy()
            if not supplement.empty:
                supplement["Bucket"] = "Supplement"
                watchlist = pd.concat([watchlist, supplement], ignore_index=True)
        watchlist = watchlist.reset_index(drop=True)
        watchlist.insert(0, "Rank", watchlist.index + 1)
        return watchlist

    def _select_watchlist_from_snapshot(self, snapshot: pd.DataFrame) -> pd.DataFrame:
        if snapshot.empty:
            return snapshot
        df = snapshot.copy()
        if "MissingCount" not in df.columns:
            df = self._add_missing_fields(df)
        df["Code"] = self._normalize_code_series(df["Code"])
        score_col = self._score_column(self.params, df)
        if "BaseScreenFlag" in df.columns:
            confirmed = df[df["BaseScreenFlag"] == True].sort_values(score_col, ascending=False)  # noqa: E712
            watchlist = confirmed.head(self.watchlist_size).copy()
            watchlist["Bucket"] = "Confirmed"
            remaining_slots = self.watchlist_size - len(watchlist)
            if remaining_slots > 0:
                unconfirmed_pool = df[df["BaseScreenFlag"].isna()].copy()
                unconfirmed_pool = unconfirmed_pool.sort_values(["MissingCount", score_col], ascending=[True, False])
                unconfirmed = unconfirmed_pool.head(remaining_slots).copy()
                if not unconfirmed.empty:
                    unconfirmed["Bucket"] = "Unconfirmed"
                    watchlist = pd.concat([watchlist, unconfirmed], ignore_index=True)
            remaining_slots = self.watchlist_size - len(watchlist)
            if remaining_slots > 0:
                supplement_pool = df[df["BaseScreenFlag"] == False].copy()  # noqa: E712
                supplement_pool = supplement_pool.sort_values(score_col, ascending=False)
                supplement = supplement_pool.head(remaining_slots).copy()
                if not supplement.empty:
                    supplement["Bucket"] = "Supplement"
                    watchlist = pd.concat([watchlist, supplement], ignore_index=True)
        elif "Bucket" in df.columns:
            watchlist = df.copy()
        else:
            watchlist = df.sort_values(score_col, ascending=False).head(self.watchlist_size).copy()
        watchlist = watchlist.reset_index(drop=True)
        watchlist.insert(0, "Rank", watchlist.index + 1)
        return watchlist

    def _build_snapshot(self, df: pd.DataFrame, run_id: str, bucket_col: str = "Bucket") -> pd.DataFrame:
        snapshot = self._add_missing_fields(df)
        snapshot["run_id"] = run_id
        if bucket_col not in snapshot.columns:
            snapshot[bucket_col] = pd.NA
        snapshot["Code"] = self._normalize_code_series(snapshot["Code"])
        required_cols = [
            "run_id",
            "Rank",
            "Code",
            "Name",
            "MarketSegment",
            "Sector33",
            "PBR_used",
            "NetCash",
            "NetCashRatio",
            "RetainedEarningsRatio",
            "ADV20Value",
            "TSE_PBR_Status",
            "TSE_PBR_EnglishFlag",
            "TSE_PBR_InvestorContactFlag",
            "EDINET_LH_Count_LB",
            "EDINET_LH_LastSubmitDate",
            "ValueScore",
            "BalanceScore",
            "LiquidityScore",
            "CatalystScore",
            "TotalScore",
            "TotalScoreAdj",
            "DataQualityPenalty",
            "MissingCount",
            "MissingFields",
            "BaseScreenFlag",
            bucket_col,
        ]
        cols = [col for col in required_cols if col in snapshot.columns]
        return snapshot[cols].copy()

    def _extract_prev_curr_runs(self, history: pd.DataFrame, curr_run_id: str) -> Tuple[pd.DataFrame, pd.DataFrame, Optional[str]]:
        if history.empty or "run_id" not in history.columns:
            return pd.DataFrame(), pd.DataFrame(), None
        history = history.copy()
        history["run_id_parsed"] = pd.to_datetime(history["run_id"], errors="coerce")
        runs = history[["run_id", "run_id_parsed"]].dropna().drop_duplicates().sort_values("run_id_parsed")
        curr_idx = runs.index[runs["run_id"] == curr_run_id].tolist()
        if not curr_idx:
            return pd.DataFrame(), pd.DataFrame(), None
        curr_pos = runs.index.get_loc(curr_idx[0])
        prev_run_id = runs.iloc[curr_pos - 1]["run_id"] if curr_pos > 0 else None
        curr_snapshot = history[history["run_id"] == curr_run_id].drop(columns=["run_id_parsed"])
        prev_snapshot = history[history["run_id"] == prev_run_id].drop(columns=["run_id_parsed"]) if prev_run_id else pd.DataFrame()
        return prev_snapshot, curr_snapshot, prev_run_id

    def _build_new_signal_table(self, merged: pd.DataFrame) -> pd.DataFrame:
        rows = []
        for _, row in merged.iterrows():
            changes = []
            if row.get("TSE_PBR_Status_prev") != "開示済" and row.get("TSE_PBR_Status_curr") == "開示済":
                changes.append("TSE_PBR_Status:開示済")
            if row.get("TSE_PBR_EnglishFlag_prev") != 1 and row.get("TSE_PBR_EnglishFlag_curr") == 1:
                changes.append("TSE_PBR_EnglishFlag:1")
            if row.get("TSE_PBR_InvestorContactFlag_prev") != 1 and row.get("TSE_PBR_InvestorContactFlag_curr") == 1:
                changes.append("TSE_PBR_InvestorContactFlag:1")
            prev_count = row.get("EDINET_LH_Count_LB_prev")
            curr_count = row.get("EDINET_LH_Count_LB_curr")
            if (pd.isna(prev_count) or prev_count < 1) and pd.notna(curr_count) and curr_count >= 1:
                changes.append("EDINET_LH_Count_LB:>=1")
            if changes:
                rows.append(
                    {
                        "Code": row.get("Code"),
                        "Name": row.get("Name_curr") or row.get("Name_prev"),
                        "NewSignal": ";".join(changes),
                        "EDINET_LH_LastSubmitDate": row.get("EDINET_LH_LastSubmitDate_curr"),
                    }
                )
        return pd.DataFrame(rows)

    def _compute_diff_data(
        self,
        prev_snapshot: pd.DataFrame,
        curr_snapshot: pd.DataFrame,
        prev_watchlist: pd.DataFrame,
        curr_watchlist: pd.DataFrame,
    ) -> Dict[str, pd.DataFrame]:
        if prev_snapshot.empty or curr_snapshot.empty:
            return {
                "diff_sheet": pd.DataFrame([["前回スナップショットが無いため差分は作成されていません。"]], columns=["Notice"]),
                "in_df": pd.DataFrame(),
                "out_df": pd.DataFrame(),
                "movers_up": pd.DataFrame(),
                "movers_down": pd.DataFrame(),
                "new_signal": pd.DataFrame(),
            }
        score_col = self._score_column(self.params, curr_snapshot)
        prev_watch_codes = set(self._normalize_code_series(prev_watchlist["Code"]))
        curr_watch_codes = set(self._normalize_code_series(curr_watchlist["Code"]))
        in_codes = sorted(curr_watch_codes - prev_watch_codes)
        out_codes = sorted(prev_watch_codes - curr_watch_codes)

        prev_cols = prev_snapshot.add_suffix("_prev")
        curr_cols = curr_snapshot.add_suffix("_curr")
        merged = curr_cols.merge(prev_cols, left_on="Code_curr", right_on="Code_prev", how="left")
        merged["Code"] = merged["Code_curr"]
        merged["TotalScore_change"] = merged[f"{score_col}_curr"] - merged[f"{score_col}_prev"]
        merged["CatalystScore_change"] = merged["CatalystScore_curr"] - merged["CatalystScore_prev"]
        merged["Rank_change"] = merged["Rank_prev"] - merged["Rank_curr"]

        in_df = merged[merged["Code"].isin(in_codes)].copy()
        out_df = prev_snapshot[prev_snapshot["Code"].isin(out_codes)].copy()
        in_df = in_df.assign(
            Name=in_df["Name_curr"],
            Rank=in_df["Rank_curr"],
            TotalScoreUsed=in_df[f"{score_col}_curr"],
            MissingFields=in_df["MissingFields_curr"],
            BaseScreenFlagChange=in_df["BaseScreenFlag_prev"].astype(str) + "→" + in_df["BaseScreenFlag_curr"].astype(str),
        )
        out_df = out_df.assign(
            Name=out_df["Name"],
            Rank=out_df["Rank"],
            TotalScoreUsed=out_df[score_col] if score_col in out_df.columns else out_df.get("TotalScore"),
            MissingFields=out_df.get("MissingFields"),
        )

        movers = merged.dropna(subset=[f"{score_col}_curr", f"{score_col}_prev"]).copy()
        movers["TotalScore_change"] = movers[f"{score_col}_curr"] - movers[f"{score_col}_prev"]
        score_threshold = float(self.params.get("SCORE_MOVE_THRESHOLD", 5))
        rank_threshold = float(self.params.get("RANK_MOVE_THRESHOLD", 10))
        movers["Rank_change"] = movers["Rank_prev"] - movers["Rank_curr"]
        focus_mask = movers["TotalScore_change"].abs().ge(score_threshold) | movers["Rank_change"].abs().ge(rank_threshold)
        movers_focus = movers[focus_mask].copy()
        movers_up = movers_focus.sort_values("TotalScore_change", ascending=False).head(10)
        movers_down = movers_focus.sort_values("TotalScore_change", ascending=True).head(10)

        new_signal = self._build_new_signal_table(merged)

        rows: List[List[Any]] = []
        rows.append(["IN (新規)"])
        rows.append(["Code", "Name", "Rank", "TotalScore(Adj)", "MissingFields", "ScoreChange", "CatalystChange", "BaseScreenFlagChange"])
        for _, row in in_df.iterrows():
            rows.append(
                [
                    row.get("Code"),
                    row.get("Name"),
                    row.get("Rank"),
                    row.get("TotalScoreUsed"),
                    row.get("MissingFields"),
                    row.get("TotalScore_change"),
                    row.get("CatalystScore_change"),
                    row.get("BaseScreenFlagChange"),
                ]
            )
        rows.append([])
        rows.append(["OUT (除外)"])
        rows.append(["Code", "Name", "Rank", "TotalScore(Adj)", "MissingFields"])
        for _, row in out_df.iterrows():
            rows.append([row.get("Code"), row.get("Name"), row.get("Rank"), row.get("TotalScoreUsed"), row.get("MissingFields")])
        rows.append([])
        rows.append(["注目変化: 上昇 Top10"])
        rows.append(["Code", "Name", "Rank_change", "TotalScore_change", "CatalystScore_change"])
        for _, row in movers_up.iterrows():
            rows.append(
                [
                    row.get("Code"),
                    row.get("Name_curr"),
                    row.get("Rank_change"),
                    row.get("TotalScore_change"),
                    row.get("CatalystScore_change"),
                ]
            )
        rows.append([])
        rows.append(["注目変化: 下落 Top10"])
        rows.append(["Code", "Name", "Rank_change", "TotalScore_change", "CatalystScore_change"])
        for _, row in movers_down.iterrows():
            rows.append(
                [
                    row.get("Code"),
                    row.get("Name_curr"),
                    row.get("Rank_change"),
                    row.get("TotalScore_change"),
                    row.get("CatalystScore_change"),
                ]
            )
        rows.append([])
        rows.append(["NewSignal"])
        rows.append(["Code", "Name", "NewSignal", "EDINET_LH_LastSubmitDate"])
        for _, row in new_signal.iterrows():
            rows.append([row.get("Code"), row.get("Name"), row.get("NewSignal"), row.get("EDINET_LH_LastSubmitDate")])

        max_cols = max(len(r) for r in rows) if rows else 1
        normalized_rows = [r + [""] * (max_cols - len(r)) for r in rows]
        columns = [f"col_{i+1}" for i in range(max_cols)]
        diff_sheet = pd.DataFrame(normalized_rows, columns=columns)
        return {
            "diff_sheet": diff_sheet,
            "in_df": in_df,
            "out_df": out_df,
            "movers_up": movers_up,
            "movers_down": movers_down,
            "movers_focus": movers_focus,
            "new_signal": new_signal,
        }

    def _confirmed_from_snapshot(self, snapshot: pd.DataFrame) -> pd.DataFrame:
        if snapshot.empty or "BaseScreenFlag" not in snapshot.columns:
            return pd.DataFrame()
        score_col = self._score_column(self.params, snapshot)
        confirmed = snapshot[snapshot["BaseScreenFlag"] == True].copy()  # noqa: E712
        confirmed = confirmed.sort_values(score_col, ascending=False).head(20).reset_index(drop=True)
        confirmed.insert(0, "Rank", confirmed.index + 1)
        return confirmed

    def _build_weekly_report(
        self,
        run_id: str,
        confirmed: pd.DataFrame,
        watchlist: pd.DataFrame,
        diff_data: Dict[str, pd.DataFrame],
        prev_exists: bool,
        scores: Optional[pd.DataFrame],
    ) -> pd.DataFrame:
        diff_sheet = diff_data.get("diff_sheet", pd.DataFrame())
        in_df = diff_data.get("in_df", pd.DataFrame())
        out_df = diff_data.get("out_df", pd.DataFrame())
        new_signal = diff_data.get("new_signal", pd.DataFrame())
        rows: List[List[Any]] = []
        rows.append(["Weekly Report"])
        rows.append(["run_id", run_id])
        rows.append(["SCORE_MODE", self.params.get("SCORE_MODE")])
        rows.append([
            "Filters",
            f"PBR_threshold={self.params.get('PBR_threshold')}, ADV20Value_min={self.params.get('ADV20Value_min')}, lookback_days={self.params.get('LOOKBACK_DAYS_EDINET')}",
        ])
        rows.append([])
        rows.append(["Summary"])
        if scores is not None and "BaseScreenFlag" in scores.columns:
            base_counts = scores["BaseScreenFlag"].value_counts(dropna=False).to_dict()
            rows.append(["- BaseScreen True/False/NA", json.dumps(base_counts, ensure_ascii=False)])
        else:
            rows.append(["- BaseScreen True/False/NA", "N/A (report-only)"])
        rows.append(["- Confirmed Top20 count", len(confirmed)])
        if prev_exists:
            rows.append(["- IN count", len(in_df)])
            rows.append(["- OUT count", len(out_df)])
        else:
            rows.append(["- IN/OUT", "前回なし"])
        movers_focus = diff_data.get("movers_focus", pd.DataFrame())
        rows.append(["- 注目変化", len(movers_focus)])
        rows.append(["- 外部シグナル新規", len(new_signal)])
        missing_top = self.audit.get("derived_missing_top")
        if missing_top:
            rows.append(["- 欠損が多い上位フィールド", f"{json.dumps(missing_top, ensure_ascii=False)} (欠損は未確認)"])
        else:
            rows.append(["- 欠損が多い上位フィールド", "N/A"])
        external_status = self._external_status() if scores is not None else "NA"
        if external_status in {"FAILED", "PARTIAL"}:
            rows.append(["- 注記", "外部取得が不完全のため新規シグナル検知は不完全です。"])
        rows.append([])

        rows.append(["Confirmed Top20"])
        rows.append(["Rank", "Code", "Name", "PBR_used", "NetCashRatio", "RetEarnRatio", "ADV20Value", "Catalyst", "Total"])
        for _, row in confirmed.iterrows():
            rows.append(
                [
                    row.get("Rank"),
                    row.get("Code"),
                    row.get("Name"),
                    row.get("PBR_used"),
                    row.get("NetCashRatio"),
                    row.get("RetainedEarningsRatio"),
                    row.get("ADV20Value"),
                    row.get("CatalystScore"),
                    row.get("TotalScoreAdj") if "TotalScoreAdj" in row.index else row.get("TotalScore"),
                ]
            )

        rows.append([])
        rows.append(["IN/OUT List"])
        rows.append(["Code", "Name", "prev_rank", "curr_rank", "TotalScore_change", "NewSignal"])
        if prev_exists:
            new_signal_codes = set(new_signal["Code"]) if not new_signal.empty else set()
            for _, row in in_df.iterrows():
                rows.append(
                    [
                        row.get("Code"),
                        row.get("Name"),
                        "",
                        row.get("Rank"),
                        row.get("TotalScore_change"),
                        "Yes" if row.get("Code") in new_signal_codes else "",
                    ]
                )
            for _, row in out_df.iterrows():
                rows.append(
                    [
                        row.get("Code"),
                        row.get("Name"),
                        row.get("Rank"),
                        "",
                        "",
                        "Yes" if row.get("Code") in new_signal_codes else "",
                    ]
                )
        else:
            rows.append(["前回なし"])

        if not new_signal.empty:
            rows.append([])
            rows.append(["NewSignal"])
            rows.append(["Code", "Name", "Signal", "EDINET_LH_LastSubmitDate"])
            for _, row in new_signal.iterrows():
                rows.append([row.get("Code"), row.get("Name"), row.get("NewSignal"), row.get("EDINET_LH_LastSubmitDate")])

        max_cols = max(len(r) for r in rows) if rows else 1
        normalized_rows = [r + [""] * (max_cols - len(r)) for r in rows]
        columns = [f"col_{i+1}" for i in range(max_cols)]
        return pd.DataFrame(normalized_rows, columns=columns)

    # ------------------------ external ------------------------
    def run_external(self):
        # TSE
        if self.no_external:
            self.external_data["TSE"] = pd.DataFrame(columns=TSEListNormalizer.REQUIRED_COLS)
            self.external_results["TSE"] = ExternalFetchResult(
                status="FAILED",
                message="Skipped by --no-external",
                fetched_at=datetime.now(),
                source="TSE",
            )
        else:
            tse = TSEListNormalizer()
            norm_df, result, audit = tse.run(
                self.params["TSE_PBR_PAGE_URL"],
                self.params.get("TSE_PBR_LIST_XLSX_URL") or None,
                self.cache_dir,
                self.refresh_cache,
            )
            self.external_data["TSE"] = norm_df
            self.external_results["TSE"] = result
            self.audit["TSE"] = audit

        # EDINET
        if self.no_external:
            self.external_data["EDINET"] = pd.DataFrame(columns=EdinetLargeHoldingExtractor.REQUIRED_COLS)
            self.external_results["EDINET"] = ExternalFetchResult(
                status="FAILED",
                message="Skipped by --no-external",
                fetched_at=datetime.now(),
                source="EDINET",
            )
        else:
            edinet = EdinetLargeHoldingExtractor()
            api_key = os.environ.get("EDINET_API_KEY", "")
            df, result, audit = edinet.run(
                int(self.params["LOOKBACK_DAYS_EDINET"]),
                self.params["EDINET_FORM_CODE_LIST_URL"],
                api_key,
                self.cache_dir,
                self.refresh_cache,
            )
            self.external_data["EDINET"] = df
            self.external_results["EDINET"] = result
            self.audit["EDINET"] = audit

    # ------------------------ derived ------------------------
    def build_derived(self, dfs: Dict[str, pd.DataFrame]) -> pd.DataFrame:
        universe = dfs["Universe"].copy()
        market = dfs["Market"].copy()
        fundamentals = dfs["Fundamentals"].copy()
        merged = universe.merge(market, on="Code", how="left", suffixes=("", "_mkt"))
        merged = merged.merge(fundamentals, on="Code", how="left")
        # Derived columns
        merged["PBR_mkt"] = merged.get("PBR")
        merged["PBR_calc"] = merged.apply(
            lambda r: r["MktCap"] / r["Equity"] if pd.notna(r.get("MktCap")) and pd.notna(r.get("Equity")) and r.get("Equity") != 0 else pd.NA,
            axis=1,
        )
        merged["PBR_used"] = merged["PBR_mkt"].combine_first(merged["PBR_calc"])
        merged["NetCash"] = merged.apply(
            lambda r: (r.get("Cash") + r.get("STInvest")) - r.get("InterestDebt") if pd.notna(r.get("Cash")) and pd.notna(r.get("InterestDebt")) and pd.notna(r.get("STInvest")) else pd.NA,
            axis=1,
        )
        merged["NetCashRatio"] = merged.apply(
            lambda r: r["NetCash"] / r["MktCap"] if pd.notna(r.get("NetCash")) and pd.notna(r.get("MktCap")) and r.get("MktCap") != 0 else pd.NA,
            axis=1,
        )
        merged["EquityRatio"] = merged.apply(
            lambda r: r.get("Equity") / r.get("TotalAssets") if pd.notna(r.get("Equity")) and pd.notna(r.get("TotalAssets")) and r.get("TotalAssets") != 0 else pd.NA,
            axis=1,
        )
        merged["RetainedEarningsRatio"] = merged.apply(
            lambda r: r.get("RetainedEarnings") / r.get("MktCap") if pd.notna(r.get("RetainedEarnings")) and pd.notna(r.get("MktCap")) and r.get("MktCap") != 0 else pd.NA,
            axis=1,
        )

        # External columns
        tse_df = self.external_data.get("TSE", pd.DataFrame(columns=TSEListNormalizer.REQUIRED_COLS))
        edinet_df = self.external_data.get("EDINET", pd.DataFrame(columns=EdinetLargeHoldingExtractor.REQUIRED_COLS))
        tse_cols = {
            "Code": "Code4",
            "TSE_PBR_Status": "Status",
            "TSE_PBR_UpdateDate": "UpdateDate",
            "TSE_PBR_EnglishFlag": "EnglishFlag",
            "TSE_PBR_InvestorContactFlag": "InvestorContactFlag",
        }
        tse_df = tse_df.rename(columns={"Code4": "Code"})
        merged = merged.merge(tse_df[["Code", "Status", "UpdateDate", "EnglishFlag", "InvestorContactFlag"]], on="Code", how="left")
        merged = merged.rename(columns={
            "Status": "TSE_PBR_Status",
            "UpdateDate": "TSE_PBR_UpdateDate",
            "EnglishFlag": "TSE_PBR_EnglishFlag",
            "InvestorContactFlag": "TSE_PBR_InvestorContactFlag",
        })

        # EDINET counts
        edinet_df = edinet_df.copy()
        if not edinet_df.empty:
            edinet_df["submitDateTime_parsed"] = pd.to_datetime(edinet_df["submitDateTime"], errors="coerce")
            edinet_active = edinet_df[edinet_df["IsActive"] == True]  # noqa: E712
            ed_summary = edinet_active.groupby("Code4").agg(
                EDINET_LH_Count_LB=("docID", "nunique"),
                EDINET_LH_LastSubmitDate=("submitDateTime_parsed", "max"),
                EDINET_LH_WeakMatchCount_LB=("IsWeakMatch", "sum"),
            ).reset_index().rename(columns={"Code4": "Code"})
        else:
            ed_summary = pd.DataFrame(columns=["Code", "EDINET_LH_Count_LB", "EDINET_LH_LastSubmitDate", "EDINET_LH_WeakMatchCount_LB"])
        merged = merged.merge(ed_summary, on="Code", how="left")

        # Flags
        merged["LiquidityFlag"] = merged["ADV20Value"].apply(
            lambda x: True if pd.notna(x) and x >= float(self.params["ADV20Value_min"]) else (pd.NA if pd.isna(x) else False)
        )
        merged["PBR1Flag"] = merged["PBR_used"].apply(
            lambda x: True if pd.notna(x) and x < float(self.params["PBR_threshold"]) else (pd.NA if pd.isna(x) else False)
        )
        merged["NetCashFlag"] = merged["NetCash"].apply(
            lambda x: True if pd.notna(x) and x > 0 else (pd.NA if pd.isna(x) else False)
        )

        def combine_flags(row):
            flags = [row["LiquidityFlag"], row["PBR1Flag"], row["NetCashFlag"]]
            if any(f is False for f in flags):
                return False
            if all(f is True for f in flags):
                return True
            return pd.NA

        merged["BaseScreenFlag"] = merged.apply(combine_flags, axis=1)
        self.audit["derived_base_counts"] = merged["BaseScreenFlag"].value_counts(dropna=False).to_dict()
        kpi_cols = ["PBR_used", "NetCash", "NetCashRatio", "RetainedEarningsRatio", "ADV20Value", "Equity", "TotalAssets"]
        missing_counts = {col: merged[col].isna().sum() for col in kpi_cols if col in merged.columns}
        missing_sorted = dict(sorted(missing_counts.items(), key=lambda x: x[1], reverse=True)[:5])
        self.audit["derived_missing_top"] = missing_sorted
        return merged

    # ------------------------ scores ------------------------
    def build_scores(self, derived: pd.DataFrame) -> pd.DataFrame:
        df = derived.copy()
        # ValueScore
        pbr_floor = float(self.params["PBR_good_floor"])
        weight_value = float(self.params["WEIGHT_VALUE"])
        def value_score(pbr):
            if pd.isna(pbr):
                return pd.NA
            if pbr <= pbr_floor:
                return weight_value
            if pbr >= float(self.params["PBR_threshold"]):
                return 0.0
            # linear between floor and threshold
            slope = weight_value / (float(self.params["PBR_threshold"]) - pbr_floor)
            return max(0.0, weight_value - slope * (pbr - pbr_floor))
        df["ValueScore"] = df["PBR_used"].apply(value_score)

        # BalanceScore
        netcap = float(self.params["NETCASHRATIO_CAP"])
        retaearncap = float(self.params["RETEARNRATIO_CAP"])
        weight_balance = float(self.params["WEIGHT_BALANCE"])
        netcash_weight = 25
        retearn_weight = 10

        def netcash_score(x):
            if pd.isna(x):
                return pd.NA
            capped = min(netcap, max(-netcap, x))
            return (capped / netcap) * netcash_weight if netcap != 0 else pd.NA

        def retearn_score(x):
            if pd.isna(x):
                return pd.NA
            capped = min(retaearncap, max(-retaearncap, x))
            return (capped / retaearncap) * retearn_weight if retaearncap != 0 else pd.NA

        df["BalanceScore"] = df["NetCashRatio"].apply(netcash_score) + df["RetainedEarningsRatio"].apply(retearn_score)

        # LiquidityScore linear
        adv_weight = float(self.params["WEIGHT_LIQUIDITY"])
        adv_min = float(self.params["ADV20Value_min"])
        def liq_score(x):
            if pd.isna(x):
                return pd.NA
            if x <= adv_min:
                return 0.0
            return min(adv_weight, adv_weight * (x / (adv_min * 5)))  # soft scale
        df["LiquidityScore"] = df["ADV20Value"].apply(liq_score)

        # CatalystScore
        def catalyst(row):
            score = 0.0
            if row.get("TSE_PBR_Status") == "開示済":
                score += float(self.params["CATALYST_TSE_DISCLOSED_POINTS"])
            if row.get("TSE_PBR_EnglishFlag") == 1:
                score += float(self.params["CATALYST_TSE_ENG_POINTS"])
            if row.get("TSE_PBR_InvestorContactFlag") == 1:
                score += float(self.params["CATALYST_TSE_CONTACT_POINTS"])
            cnt = row.get("EDINET_LH_Count_LB")
            if pd.notna(cnt) and cnt >= 1:
                base = float(self.params["CATALYST_EDINET_NEW_POINTS"])
                mult = float(self.params["CATALYST_EDINET_MULTIPLIER"])
                score += min(base, 3 + cnt * mult)
            return min(float(self.params["CATALYST_MAX"]), score)

        df["CatalystScore"] = df.apply(catalyst, axis=1)

        df["ExternalDataStatus"] = self._external_status()

        # TotalScore
        df["TotalScore"] = df["ValueScore"] + df["BalanceScore"] + df["LiquidityScore"] + df["CatalystScore"]
        if self.params.get("SCORE_MODE", "STRICT") == "STRICT":
            mask_na = df[["ValueScore", "BalanceScore", "LiquidityScore"]].isna().any(axis=1)
            df.loc[mask_na, "TotalScore"] = pd.NA
        else:
            df["TotalScore"] = (
                df["ValueScore"].fillna(0)
                + df["BalanceScore"].fillna(0)
                + df["LiquidityScore"].fillna(0)
                + df["CatalystScore"].fillna(0)
            )
            missing_counts = df[["PBR_used", "NetCashRatio", "RetainedEarningsRatio", "ADV20Value"]].isna().sum(axis=1)
            penalty = -(missing_counts * 5).clip(upper=20)
            df["DataQualityPenalty"] = penalty
            df["TotalScoreAdj"] = (df["TotalScore"].fillna(0) + penalty).clip(lower=0)
            df["MissingCount"] = missing_counts
            self.audit["missing_count_distribution"] = missing_counts.value_counts(dropna=False).to_dict()
        self.audit["scores_total_na"] = df["TotalScore"].isna().sum()
        return df

    def _external_status(self) -> str:
        statuses = [self.external_results.get("TSE"), self.external_results.get("EDINET")]
        failed = [r for r in statuses if r and r.status == "FAILED"]
        if len(failed) == 2:
            return "FAILED"
        if len(failed) == 1:
            return "PARTIAL"
        return "OK"

    # ------------------------ candidates ------------------------
    def build_candidates(self, scores: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
        df = scores.copy()
        strict_mode = self.params.get("SCORE_MODE", "STRICT") == "STRICT"
        # Confirmed
        confirmed = df[df["BaseScreenFlag"] == True]  # noqa: E712
        if strict_mode:
            confirmed = confirmed[confirmed["TotalScore"].notna()]
            confirmed = confirmed.sort_values("TotalScore", ascending=False).head(20)
        else:
            confirmed = confirmed.sort_values("TotalScoreAdj", ascending=False).head(20)
        confirmed = confirmed.reset_index(drop=True)
        confirmed.insert(0, "Rank", confirmed.index + 1)

        # Unconfirmed
        if strict_mode:
            unconf = df[(df["BaseScreenFlag"].isna()) | ((df["BaseScreenFlag"] == True) & (df["TotalScore"].isna()))]  # noqa: E712
            unconf = unconf.sort_values("PBR_used", ascending=True)
        else:
            unconf = df[df["BaseScreenFlag"].isna()]
            unconf = unconf.sort_values("TotalScoreAdj", ascending=False)
        unconf = unconf.head(20).reset_index(drop=True)
        missing_fields = []
        for _, row in unconf.iterrows():
            missing = []
            for col in ["PBR_used", "NetCashRatio", "RetainedEarningsRatio", "ADV20Value"]:
                if pd.isna(row.get(col)):
                    missing.append(col)
            missing_fields.append(",".join(missing) if missing else "")
        unconf["MissingFields"] = missing_fields
        unconf.insert(0, "Rank", unconf.index + 1)
        return confirmed, unconf

    # ------------------------ output writing ------------------------
    def write_outputs(
        self,
        derived: Optional[pd.DataFrame],
        scores: Optional[pd.DataFrame],
        confirmed: Optional[pd.DataFrame],
        unconf: Optional[pd.DataFrame],
        watchlist: pd.DataFrame,
        snapshot: pd.DataFrame,
        history: pd.DataFrame,
        diff_sheet: pd.DataFrame,
        weekly_report: pd.DataFrame,
    ):
        # ensure outdir exists
        self.outdir.mkdir(parents=True, exist_ok=True)
        out_name = self.workbook.stem + "_out.xlsx"
        out_path = self.outdir / out_name
        # replace sheets
        self.io.replace_sheet_with_df("Params", pd.DataFrame(list(self.params.items()), columns=["Param", "Value"]))
        if derived is not None:
            self.io.replace_sheet_with_df("Derived", derived)
        if scores is not None:
            self.io.replace_sheet_with_df("Scores", scores)
        if confirmed is not None and unconf is not None:
            candidates_combined = pd.concat(
                [confirmed.assign(Category="Confirmed"), unconf.assign(Category="Unconfirmed")], ignore_index=True
            )
            self.io.replace_sheet_with_df("Candidates", candidates_combined)
        if self.external_data:
            self.io.replace_sheet_with_df(
                "TSE_PBR_List", self.external_data.get("TSE", pd.DataFrame(columns=TSEListNormalizer.REQUIRED_COLS))
            )
            self.io.replace_sheet_with_df(
                "EDINET_LargeHolding",
                self.external_data.get("EDINET", pd.DataFrame(columns=EdinetLargeHoldingExtractor.REQUIRED_COLS)),
            )
        self.io.replace_sheet_with_df("Watchlist", watchlist)
        self.io.replace_sheet_with_df("Candidates_Snapshot", snapshot)
        self.io.replace_sheet_with_df("Snapshot_History", history)
        self.io.replace_sheet_with_df("Diff", diff_sheet)
        self.io.replace_sheet_with_df("Weekly_Report", weekly_report)
        self.write_audit()
        self.io.save(out_path)
        self.out_path = out_path

    def write_audit(self):
        audit_rows = []
        audit_rows.append(["run_at", self.audit.get("run_at")])
        audit_rows.append(["input_file", str(self.workbook)])
        audit_rows.append(["python_version", sys.version])
        audit_rows.append(["pandas_version", pd.__version__])
        audit_rows.append(["external_TSE_status", getattr(self.external_results.get("TSE"), "status", "NA")])
        audit_rows.append(["external_EDINET_status", getattr(self.external_results.get("EDINET"), "status", "NA")])
        audit_rows.append(["BaseScreen_counts", json.dumps(self.audit.get("derived_base_counts", {}), ensure_ascii=False)])
        audit_rows.append(["scores_total_na", self.audit.get("scores_total_na")])
        audit_rows.append(["derived_missing_top", json.dumps(self.audit.get("derived_missing_top", {}), ensure_ascii=False)])
        # Column resolutions
        audit_rows.append(["column_resolution", json.dumps(self.audit.get("column_resolution", []), ensure_ascii=False)])
        # External audits
        audit_rows.append(["TSE_audit", json.dumps(self.audit.get("TSE", {}), ensure_ascii=False)])
        audit_rows.append(["EDINET_audit", json.dumps(self.audit.get("EDINET", {}), ensure_ascii=False)])
        df = pd.DataFrame(audit_rows, columns=["item", "value"])
        self.io.replace_sheet_with_df("Audit", df)

    # ------------------------ run ------------------------
    def run(self):
        logging.info("Loading params")
        self.load_params()
        out_name = self.workbook.stem + "_out.xlsx"
        out_path = self.outdir / out_name
        history = self._load_snapshot_history(out_path)
        if self.report_only:
            if history.empty:
                run_id = self._run_id_now()
                watchlist = pd.DataFrame()
                snapshot = pd.DataFrame()
                diff_data = {"diff_sheet": pd.DataFrame([["前回なし"]], columns=["Notice"]), "new_signal": pd.DataFrame()}
                weekly_report = self._build_weekly_report(
                    run_id=run_id,
                    confirmed=pd.DataFrame(),
                    watchlist=watchlist,
                    diff_data=diff_data,
                    prev_exists=False,
                    scores=None,
                )
                self.write_outputs(None, None, None, None, watchlist, snapshot, history, diff_data["diff_sheet"], weekly_report)
                return
            history["run_id_parsed"] = pd.to_datetime(history["run_id"], errors="coerce")
            latest_run = history.dropna(subset=["run_id_parsed"]).sort_values("run_id_parsed").iloc[-1]["run_id"]
            prev_snapshot, curr_snapshot, prev_run_id = self._extract_prev_curr_runs(history, latest_run)
            watchlist = self._select_watchlist_from_snapshot(curr_snapshot)
            prev_watchlist = self._select_watchlist_from_snapshot(prev_snapshot) if not prev_snapshot.empty else pd.DataFrame()
            diff_data = self._compute_diff_data(prev_snapshot, curr_snapshot, prev_watchlist, watchlist)
            confirmed_snapshot = self._confirmed_from_snapshot(curr_snapshot)
            weekly_report = self._build_weekly_report(
                run_id=latest_run,
                confirmed=confirmed_snapshot,
                watchlist=watchlist,
                diff_data=diff_data,
                prev_exists=prev_run_id is not None,
                scores=None,
            )
            history = history.drop(columns=["run_id_parsed"])
            self.write_outputs(None, None, None, None, watchlist, curr_snapshot, history, diff_data["diff_sheet"], weekly_report)
            return

        logging.info("Loading inputs")
        dfs = self.load_inputs()
        logging.info("Running external fetches")
        self.run_external()
        logging.info("Building derived")
        derived = self.build_derived(dfs)
        logging.info("Building scores")
        scores = self.build_scores(derived)
        logging.info("Building candidates")
        confirmed, unconf = self.build_candidates(scores)
        watchlist = self._select_watchlist(scores)
        run_id = self._run_id_now()
        if self.snapshot_bucket_mode.upper() == "WATCHLIST":
            snapshot_source = watchlist.copy()
        else:
            snapshot_source = pd.concat(
                [
                    confirmed.assign(Bucket="Confirmed"),
                    unconf.assign(Bucket="Unconfirmed"),
                ],
                ignore_index=True,
            )
        snapshot = self._build_snapshot(snapshot_source, run_id=run_id)
        history = pd.concat([history, snapshot], ignore_index=True) if not history.empty else snapshot.copy()
        history = self._trim_snapshot_history(history)
        prev_snapshot, curr_snapshot, prev_run_id = self._extract_prev_curr_runs(history, run_id)
        prev_watchlist = self._select_watchlist_from_snapshot(prev_snapshot) if not prev_snapshot.empty else pd.DataFrame()
        diff_data = self._compute_diff_data(prev_snapshot, curr_snapshot, prev_watchlist, watchlist)
        weekly_report = self._build_weekly_report(
            run_id=run_id,
            confirmed=confirmed,
            watchlist=watchlist,
            diff_data=diff_data,
            prev_exists=prev_run_id is not None,
            scores=scores,
        )
        logging.info("Writing outputs")
        self.write_outputs(derived, scores, confirmed, unconf, watchlist, snapshot, history, diff_data["diff_sheet"], weekly_report)

        # stdout summaries
        base_counts = derived["BaseScreenFlag"].value_counts(dropna=False).to_dict()
        print(f"BaseScreen counts: {base_counts}")
        print("Confirmed Candidates Top20 Code:", ",".join(confirmed["Code"].head(20).astype(str)))
        print("Unconfirmed Top20 Code:", ",".join(unconf["Code"].head(20).astype(str)))
        # External status summary
        for src in ["TSE", "EDINET"]:
            res = self.external_results.get(src)
            status = res.status if res else "NA"
            print(f"External {src} status: {status}")


# --------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Build screening pipeline outputs")
    parser.add_argument("--workbook", required=True, help="Path to input workbook")
    parser.add_argument("--outdir", default="./out", help="Output directory")
    parser.add_argument("--no-external", action="store_true", help="Skip external fetches")
    parser.add_argument("--refresh-cache", action="store_true", help="Refresh cache")
    parser.add_argument("--watchlist-size", type=int, default=30, help="Watchlist size (default 30)")
    parser.add_argument("--max-history-runs", type=int, default=52, help="Max snapshot history runs to keep")
    parser.add_argument("--report-only", action="store_true", help="Regenerate report/diff from Snapshot_History")
    parser.add_argument(
        "--snapshot-bucket-mode",
        default="CANDIDATES",
        choices=["WATCHLIST", "CANDIDATES"],
        help="Snapshot scope: WATCHLIST or CANDIDATES",
    )
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR", "CRITICAL"], help="Log level")
    args = parser.parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(asctime)s %(levelname)s %(message)s")

    pipeline = ScreeningPipeline(
        workbook=Path(args.workbook),
        outdir=Path(args.outdir),
        no_external=args.no_external,
        refresh_cache=args.refresh_cache,
        log_level=args.log_level,
        watchlist_size=args.watchlist_size,
        max_history_runs=args.max_history_runs,
        report_only=args.report_only,
        snapshot_bucket_mode=args.snapshot_bucket_mode,
    )
    pipeline.run()


if __name__ == "__main__":
    main()
